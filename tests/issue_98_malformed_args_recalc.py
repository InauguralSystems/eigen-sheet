#!/usr/bin/env python3
"""Issue #98: one malformed formula must not abort or hang the whole recalc.

This covers the malformed-ARGUMENT class: every dispatcher function called
with each shape in SHAPES. It is not the whole of #78's "every function with
error-provoking arguments recalcs" item. Three known classes are left out,
each filed with its own Done-when:

  * #95  numbers outside a financial function's domain (=PMT(0,0,100),
         =NPER(1,5,5,5)) still abort recalc;
  * #96  a missing operand under an error-capturing function becomes a
         value (=ISERROR(*) is TRUE), so the (*) shapes here only assert
         "no abort";
  * #97  XLOOKUP with a return array of another size or shape returns a
         value.

The function population is DERIVED from the formula dispatcher (every
`fname == "X"` in sheet.eigs), not hand-listed, and pinned by a floor so it
cannot shrink silently. Each function is called with every argument shape in
SHAPES, all in one sheet between two sentinel cells: C1 = A1+41 is written
before the cases and C2 = A2+40 after them. The gate is:

  * the interpreter exits 0 and both sentinels compute 42, i.e. no formula
    aborted (or hung) the recalc of the whole sheet, and recalc reached the
    cells written last;
  * every cell displays a non-empty value, and every incomplete-range shape
    displays an error token (Calc rejects a dangling ':' as Err:511 whatever
    the function, IFERROR included);
  * SCALAR_FOR_RANGE pins the result class of a scalar passed where a range is
    read. Functions that need a reference error, as Calc does (Err:504), and
    the ones Calc treats as a 1x1 array return its value. MISSING_REQUIRED
    pins that a left-out required argument is an error (Calc: Err:511).
    DOMAIN_GUARDS pins the argument values that used to abort recalc from
    inside one function (CHAR's NUL, XLOOKUP's short return array).

`--calc` additionally runs the incomplete-range cells and the tables through
headless LibreOffice and requires Calc to agree with every expected class,
except the rows named in NOT_CALC_COMPARABLE, each with its reason. It never
skips: with no LibreOffice it fails (the calc-oracle CI job installs it).
"""
import csv
import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile

EIGS = os.environ.get("EIGENSCRIPT", "eigenscript")
SOFFICE = os.environ.get("SOFFICE", "soffice")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Count of distinct dispatcher names when this gate landed. Adding functions
# is fine; a derivation that finds fewer means the population (or this
# derivation) broke, which must not pass as "every function recalcs".
FUNC_FLOOR = 115

# (label, formula template, must display an error token)
SHAPES = [
    ("incomplete range", "=%s(A1:)", True),
    ("incomplete range, later argument", "=%s(A1:A3,A1:)", True),
    ("incomplete range, middle argument", "=%s(1,A1:,1)", True),
    ("empty argument list", "=%s()", False),
    ("missing operand", "=%s(*)", False),
    ("missing operand, later argument", "=%s(1,*)", False),
    ("text for number", '=%s("abc")', False),
    ("text cells for numbers", "=%s(A4,A4,A4)", False),
    ("scalars for ranges", "=%s(5,5,5)", False),
]

# formula -> "error" or the exact display. Checked against Calc by --calc.
SCALAR_FOR_RANGE = [
    ("=ROWS(5)", "error"),
    ("=COLUMNS(5)", "error"),
    ("=ROW(5)", "error"),
    ("=COLUMN(5)", "error"),
    ("=INDEX(5,1,1)", "error"),
    ("=MATCH(5,5,0)", "error"),
    ("=VLOOKUP(5,5,1)", "error"),
    ("=HLOOKUP(5,5,1)", "error"),
    ("=IRR(5)", "error"),
    ("=SUMIF(5,5,5)", "error"),
    ("=SUMIF(A1:A3,1,5)", "error"),
    ("=COUNTIF(A1:A3,1,5)", "error"),
    ("=COUNTIF(5,5)", "error"),
    ("=AVERAGEIF(5,5,5)", "error"),
    ("=SUMIFS(5,5,5)", "error"),
    ("=COUNTIFS(5,5)", "error"),
    ("=AVERAGEIFS(5,5,5)", "error"),
    ('=CHAR("abc")', "error"),
    ("=PERCENTILE(5,0.5)", "5"),
    ("=RANK(5,5)", "1"),
    ("=SUMPRODUCT(5)", "5"),
    ("=LOOKUP(5,5,5)", "5"),
]

# A required argument left out is an error (Calc: Err:511), not a 0 default
# that the annuity math then divides by. Checked against Calc by --calc.
MISSING_REQUIRED = ["=PMT(0.1,10)", "=FV(0.1,10)", "=PV(0.1,10)",
                    "=NPER(0.1,10)", "=RATE(10,100)", "=IPMT(0.1,1,10)",
                    "=PPMT(0.1,1,10)", "=NPV(0.1)", "=COUNTIFS()",
                    "=SUMIFS(A1:A3,A1:A3)", "=SUMPRODUCT()"]

# An argument value that used to abort recalc from inside one function's
# branch, past the generic shapes above. Each must display an error token.
DOMAIN_GUARDS = ["=CHAR(0)", "=CHAR(256)", "=XLOOKUP(3,A1:A3,B1:B2)"]

# Expected-class rows --calc does not compare, with why. Every row here is
# still gated on the eigen-sheet side (it must display an error token).
NOT_CALC_COMPARABLE = {
    "=CHAR(0)": "Calc returns a NUL character; eigen-sheet strings cannot "
                "hold one, so #VALUE! (Excel's answer) is a deliberate divergence",
    "=XLOOKUP(3,A1:A3,B1:B2)": "LibreOffice 24.2 (the CI image) has no XLOOKUP "
                               "(#NAME?), so agreement would be trivial (#97)",
}

# Shape cells whose correct display IS empty. Every other cell must display
# something non-empty, so a cell recalc never reached cannot pass as a value.
# These are pinned to exactly "" and --calc requires Calc to show "" too.
EMPTY_RESULTS = ["=CONCATENATE()",   # joins nothing
                 "=MID(5,5,5)"]      # starts past the end of "5"

FIXTURE = [("A1", "1"), ("A2", "2"), ("A3", "3"), ("A4", "abc"),
           ("B1", "10"), ("B2", "20"), ("B3", "30")]
SENTINEL = ("C1", "=A1+41", "42")        # written before every case
TAIL_SENTINEL = ("C2", "=A2+40", "42")   # written after every case

# Calc stores these post-OOXML-1.0 functions with an _xlfn. prefix.
_XLFN = ("XOR", "IFS", "SWITCH", "IFNA", "XLOOKUP")


def derive_functions():
    src = open(os.path.join(ROOT, "sheet.eigs"), encoding="utf-8").read()
    return sorted(set(re.findall(r'fname == "([A-Z][A-Z0-9_.]*)"', src)))


def col_name(n):
    out = ""
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def quote(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'


def build_cases(funcs):
    """[(addr, formula, label, must_error, expected_display_or_None)]"""
    cases = []
    for fi, fn in enumerate(funcs):
        for si, (label, tmpl, must_error) in enumerate(SHAPES):
            addr = "%s%d" % (col_name(4 + si), fi + 1)   # columns D..
            formula = tmpl % fn
            want = "" if formula in EMPTY_RESULTS else None
            cases.append((addr, formula, label, must_error, want))
    col = col_name(4 + len(SHAPES))
    for i, (formula, want) in enumerate(SCALAR_FOR_RANGE, 1):
        cases.append(("%s%d" % (col, i), formula, "scalar for range",
                      want == "error", None if want == "error" else want))
    col = col_name(5 + len(SHAPES))
    for i, formula in enumerate(MISSING_REQUIRED, 1):
        cases.append(("%s%d" % (col, i), formula, "missing required argument",
                      True, None))
    col = col_name(6 + len(SHAPES))
    for i, formula in enumerate(DOMAIN_GUARDS, 1):
        cases.append(("%s%d" % (col, i), formula, "argument outside the domain",
                      True, None))
    return cases


def run_sheet(tmp, cases, timeout):
    """Recalc one sheet holding `cases`; returns (rc, {addr: display}, stderr).
    rc is None on timeout (a formula that never finishes evaluating)."""
    lines = ["import sheet", "s is sheet.new_sheet of null"]
    for addr, raw in FIXTURE + [SENTINEL[:2]]:
        lines.append("sheet.set_cell of [s, %s, %s]" % (quote(addr), quote(raw)))
    for addr, formula, _, _, _ in cases:
        lines.append("sheet.set_cell of [s, %s, %s]" % (quote(addr), quote(formula)))
    lines.append("sheet.set_cell of [s, %s, %s]"
                 % (quote(TAIL_SENTINEL[0]), quote(TAIL_SENTINEL[1])))
    lines.append("sheet.recalc of s")
    for addr in [SENTINEL[0], TAIL_SENTINEL[0]] + [c[0] for c in cases]:
        lines.append('print of (%s + "\\t" + (sheet.display of [s, %s]))'
                     % (quote(addr), quote(addr)))
    app = os.path.join(tmp, "app.eigs")
    with open(app, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    try:
        run = subprocess.run([EIGS, app], cwd=tmp, text=True,
                             capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        return None, {}, "timed out after %ss" % timeout
    got = dict(line.split("\t", 1) for line in run.stdout.splitlines()
               if "\t" in line)
    return run.returncode, got, run.stderr


def sheet_completed(rc, got):
    return (rc == 0 and got.get(SENTINEL[0]) == SENTINEL[2]
            and got.get(TAIL_SENTINEL[0]) == TAIL_SENTINEL[2])


def name_aborts(tmp, cases):
    """The whole-sheet recalc aborted: find each formula that does it alone."""
    culprits = []
    for case in cases:
        rc, got, err = run_sheet(tmp, [case], 10)
        if not sheet_completed(rc, got):
            last = (err.strip().splitlines() or ["rc=%s" % rc])[0]
            culprits.append((case, last))
    return culprits


def calc_values(tmp, formulas):
    """LibreOffice's displayed result for each formula (fixture in A/B)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for addr, raw in FIXTURE:
        ws[addr] = float(raw) if raw.replace(".", "").isdigit() else raw
    for i, formula in enumerate(formulas, 1):
        for name in _XLFN:
            formula = re.sub(r"\b" + name + r"\(", "_xlfn." + name + "(", formula)
        ws.cell(i, 26).value = formula            # column Z, clear of the fixture
    wb.calculation.fullCalcOnLoad = True
    xlsx = os.path.join(tmp, "issue98_calc.xlsx")
    wb.save(xlsx)
    profile = "file://" + os.path.join(tmp, "lo-profile")
    subprocess.run([SOFFICE, "--headless", "--norestore", "--nolockcheck",
                    "-env:UserInstallation=" + profile, "--convert-to", "csv",
                    "--outdir", tmp, xlsx], capture_output=True, timeout=300)
    csvs = glob.glob(os.path.join(tmp, "issue98_calc.csv"))
    if not csvs:
        raise RuntimeError("LibreOffice produced no CSV")
    rows = list(csv.reader(open(csvs[0], newline="", encoding="utf-8")))
    return [rows[i][25] if i < len(rows) and len(rows[i]) > 25 else ""
            for i in range(len(formulas))]


def calc_is_error(value):
    return value.startswith("#") or value.startswith("Err:")


def main(argv):
    with_calc = "--calc" in argv
    failures = 0
    funcs = derive_functions()
    if len(funcs) < FUNC_FLOOR:
        print("FAIL: derived %d dispatcher functions, floor is %d"
              % (len(funcs), FUNC_FLOOR))
        failures += 1
    else:
        print("PASS: derived %d dispatcher functions (floor %d)"
              % (len(funcs), FUNC_FLOOR))
    cases = build_cases(funcs)

    tmp = tempfile.mkdtemp(prefix="eigen-sheet-issue-98-")
    try:
        mod = os.path.join(tmp, "eigs_modules", "sheet")
        os.makedirs(mod)
        shutil.copy(os.path.join(ROOT, "sheet.eigs"), mod)
        shutil.copy(os.path.join(ROOT, "eigs.json"), mod)

        rc, got, err = run_sheet(tmp, cases, 60)
        if not sheet_completed(rc, got):
            print("FAIL: recalc of the %d-formula sheet did not complete "
                  "(rc=%s, sentinels %s=%r %s=%r): %s"
                  % (len(cases), rc, SENTINEL[0], got.get(SENTINEL[0]),
                     TAIL_SENTINEL[0], got.get(TAIL_SENTINEL[0]),
                     (err.strip().splitlines() or [""])[0]))
            culprits = name_aborts(tmp, cases)
            for (addr, formula, label, _, _), why in culprits:
                print("FAIL: %s (%s) aborts or hangs the whole recalc: %s"
                      % (formula, label, why))
            print("FAIL: %d formula(s) abort or hang recalc" % len(culprits))
            return 1
        print("PASS: %d formulas over %d shapes recalced, sentinels %s=%s %s=%s"
              % (len(cases), len(SHAPES), SENTINEL[0], SENTINEL[2],
                 TAIL_SENTINEL[0], TAIL_SENTINEL[2]))

        checked = 0
        stale = sorted(set(EMPTY_RESULTS) - set(c[1] for c in cases))
        if stale:
            failures += 1
            print("FAIL: EMPTY_RESULTS names a formula that is not a case: %s"
                  % stale)
        for addr, formula, label, must_error, want in cases:
            shown = got.get(addr)
            if shown is None or (shown == "" and want != ""):
                failures += 1
                print("FAIL: %s (%s) displayed nothing (%r)"
                      % (formula, label, shown))
            elif must_error and not shown.startswith("#"):
                failures += 1
                print("FAIL: %s (%s) displayed %r, expected an error token"
                      % (formula, label, shown))
            elif want is not None and shown != want:
                failures += 1
                print("FAIL: %s (%s) displayed %r, expected %r"
                      % (formula, label, shown, want))
            else:
                checked += 1
        print("PASS: %d/%d cells displayed their expected class"
              % (checked, len(cases)))

        if with_calc:
            if shutil.which(SOFFICE) is None:
                print("FAIL: --calc needs LibreOffice (%s) and it is missing" % SOFFICE)
                return 1
            # every error-class cell and every pinned value, in one conversion
            calc_cases = [c for c in cases if (c[3] or c[4] is not None)
                          and c[1] not in NOT_CALC_COMPARABLE]
            uncompared = [c[1] for c in cases if c[1] in NOT_CALC_COMPARABLE]
            if sorted(uncompared) != sorted(NOT_CALC_COMPARABLE):
                failures += 1
                print("FAIL: NOT_CALC_COMPARABLE names a row that is not a case: %s"
                      % sorted(set(NOT_CALC_COMPARABLE) - set(uncompared)))
            for formula in uncompared:
                print("NOTE: %s not compared with Calc: %s"
                      % (formula, NOT_CALC_COMPARABLE[formula]))
            calc = calc_values(tmp, [c[1] for c in calc_cases])
            agree = 0
            for (addr, formula, label, must_error, want), lo in zip(calc_cases, calc):
                ok = calc_is_error(lo) if must_error else lo == want
                if ok:
                    agree += 1
                else:
                    failures += 1
                    print("FAIL: Calc disagrees on %s (%s): libreoffice=%r, "
                          "expected %s" % (formula, label, lo,
                                           "an error" if must_error else repr(want)))
            print("%s: LibreOffice agrees on %d/%d expected classes"
                  % ("PASS" if agree == len(calc_cases) else "FAIL",
                     agree, len(calc_cases)))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    if failures:
        print("FAIL: %d issue #98 malformed-argument check(s) failed" % failures)
        return 1
    print("All issue #98 malformed-argument recalc checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
